import eventlet
eventlet.monkey_patch()
# -*- coding: utf-8 -*-
import os
import socket
import re
import time
import random
import copy
import zipfile
from io import BytesIO
import xml.etree.ElementTree as ET

import cloudinary
import cloudinary.uploader
from docx import Document
from flask import Flask, render_template, request, jsonify, send_file, redirect
from flask_socketio import SocketIO, emit, join_room, leave_room
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- CẤU HÌNH CLOUDINARY ---
cloudinary.config(
    cloud_name = "y0xsqdev",
    api_key = "974245233197575",
    api_secret = "tB1Y4FYGtYrv5We4OoeK8z-9-A0",
    secure = True
)

def upload_bytes_to_cloudinary(image_bytes, filename="image.png"):
    """Tải trực tiếp byte ảnh lên Cloudinary và trả về URL ảnh đã tối ưu"""
    try:
        res = cloudinary.uploader.upload(
            image_bytes,
            folder="thithi_lan",
            transformation=[{'quality': 'auto', 'fetch_format': 'auto'}]
        )
        return res.get('secure_url', '')
    except Exception as e:
        print(f"Lỗi tải ảnh Cloudinary ({filename}): {e}")
        return ""

def extract_docx_with_cloudinary(docx_path):
    """
    Bóc tách toàn bộ văn bản và ảnh từ file Word .docx một cách tuyệt đối:
    1. Đọc word/media/ bằng zipfile để đẩy toàn bộ ảnh lên Cloudinary.
    2. Đọc word/_rels/document.xml.rels để ánh xạ rId sang tên ảnh.
    3. Đọc tuần tự từng paragraph trong word/document.xml để chèn đúng vị trí ảnh.
    """
    try:
        with zipfile.ZipFile(docx_path, 'r') as zf:
            file_list = zf.namelist()

            # 1. Tải toàn bộ ảnh trong word/media/ lên Cloudinary
            media_urls = {} # media_filename -> cloudinary_url
            for fname in file_list:
                if fname.startswith('word/media/'):
                    short_name = os.path.basename(fname)
                    img_data = zf.read(fname)
                    c_url = upload_bytes_to_cloudinary(img_data, short_name)
                    if c_url:
                        media_urls[short_name] = c_url

            # 2. Đọc quan hệ rId trong document.xml.rels
            rel_to_url = {} # rId -> cloudinary_url
            rels_path = 'word/_rels/document.xml.rels'
            if rels_path in file_list and media_urls:
                rels_xml = zf.read(rels_path)
                root_rels = ET.fromstring(rels_xml)
                for rel in root_rels:
                    r_id = rel.get('Id')
                    target = rel.get('Target', '')
                    target_name = os.path.basename(target)
                    if target_name in media_urls:
                        rel_to_url[r_id] = media_urls[target_name]

            # 3. Phân tích nội dung tuần tự từ document.xml
            doc_xml = zf.read('word/document.xml')
            root_doc = ET.fromstring(doc_xml)

            # Các namespace chuẩn của file OpenXML Word
            ns = {
                'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                'v': 'urn:schemas-microsoft-com:vml'
            }

            paragraphs = []
            for p in root_doc.findall('.//w:p', ns):
                p_text_parts = []
                for elem in p.iter():
                    tag = elem.tag
                    # Thẻ chứa chữ
                    if tag == f"{{{ns['w']}}}t" and elem.text:
                        p_text_parts.append(elem.text)
                    # Thẻ chứa ảnh (DrawingML blip)
                    elif tag == f"{{{ns['a']}}}blip":
                        embed_id = elem.get(f"{{{ns['r']}}}embed")
                        if embed_id and embed_id in rel_to_url:
                            img_link = rel_to_url[embed_id]
                            p_text_parts.append(f'<br><img src="{img_link}" class="exam-img" style="max-width:100%; height:auto; margin:8px 0; display:block;" /><br>')
                    # Thẻ chứa ảnh dạng VML (ảnh cũ hoặc copy từ trình duyệt)
                    elif tag == f"{{{ns['v']}}}imagedata":
                        rel_id = elem.get(f"{{{ns['r']}}}id")
                        if rel_id and rel_id in rel_to_url:
                            img_link = rel_to_url[rel_id]
                            p_text_parts.append(f'<br><img src="{img_link}" class="exam-img" style="max-width:100%; height:auto; margin:8px 0; display:block;" /><br>')

                p_full = "".join(p_text_parts).strip()
                if p_full:
                    for sub in p_full.split('\n'):
                        s = sub.strip()
                        if s:
                            paragraphs.append(s)

            return paragraphs
    except Exception as e:
        print(f"Lỗi bóc tách docx bằng zipfile: {e}")
        doc = Document(docx_path)
        paragraphs = []
        for p in doc.paragraphs:
            for sub in p.text.split('\n'):
                s = sub.strip()
                if s:
                    paragraphs.append(s)
        return paragraphs

app = Flask(__name__)
app.config['SECRET_KEY'] = 'lan_exam_secret_key_2026'
socketio = SocketIO(app, cors_allowed_origins="*")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SUBMISSION_ROOT = os.path.join(BASE_DIR, 'bailamthisinh')
UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')

for d in [SUBMISSION_ROOT, UPLOAD_DIR]:
    if not os.path.exists(d):
        os.makedirs(d)

rooms = {}

def get_or_create_room(room_id):
    room_id = str(room_id).strip() if room_id else "default"
    if room_id not in rooms:
        room_dir = os.path.join(SUBMISSION_ROOT, room_id)
        if not os.path.exists(room_dir):
            os.makedirs(room_dir)
        rooms[room_id] = {
            "room_id": room_id,
            "class_name": "",
            "raw_bank": {"filename": "", "mcq": [], "tf": [], "sa": []},
            "active_exam": {
                "num_mcq": 0,
                "score_mcq": 0.0,
                "score_per_mcq": 0.0,
                "num_tf": 0,
                "score_tf": 0.0,
                "score_per_tf": 0.0,
                "num_sa": 0,
                "score_sa": 0.0,
                "score_per_sa": 0.0,
                "duration": 15,
                "tf_scale": {"1": 0.1, "2": 0.25, "3": 0.5, "4": 1.0},
                "status": "waiting",
                "start_time": None
            },
            "students": {},
            "submitted_ips": set()
        }
    return rooms[room_id]

def get_submission_dir(room_id):
    d = os.path.join(SUBMISSION_ROOT, str(room_id))
    if not os.path.exists(d):
        os.makedirs(d)
    return d

def clear_submission_folder(room_id):
    room_dir = get_submission_dir(room_id)
    if os.path.exists(room_dir):
        for fname in os.listdir(room_dir):
            fpath = os.path.join(room_dir, fname)
            try:
                if os.path.isfile(fpath):
                    os.remove(fpath)
            except Exception:
                pass

def auto_exam_watcher():
    while True:
        eventlet.sleep(1)
        for room_id, rdata in list(rooms.items()):
            active = rdata["active_exam"]
            if active["status"] == "running" and active["start_time"]:
                elapsed = time.time() - active["start_time"]
                total_duration = active["duration"] * 60

                if elapsed >= total_duration + 3:
                    has_updates = False
                    for sid, st in list(rdata["students"].items()):
                        if st.get("submitted") != "Đã nộp":
                            st["time_left"] = 0
                            process_grading(room_id, sid, st.get("answers", {}), st.get("name"))
                            has_updates = True

                    if has_updates:
                        active["status"] = "finished"
                        socketio.emit('update_teacher_list', list(rdata["students"].values()), to=room_id)

eventlet.spawn(auto_exam_watcher)

def get_client_ip(req):
    if req.headers.get('X-Forwarded-For'):
        return req.headers.get('X-Forwarded-For').split(',')[0].strip()
    return req.remote_addr or '127.0.0.1'

def get_base_url():
    render_url = os.environ.get('RENDER_EXTERNAL_URL')
    if render_url:
        return render_url.rstrip('/')
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
    except Exception:
        ip = '127.0.0.1'
    finally:
        s.close()
    return f"http://{ip}:5000"

def clean_question_text(raw_text):
    if not raw_text:
        return ""
    return re.sub(r'^(câu\s*\d+[\s\:\.\-\)]*)\s*', '', raw_text.strip(), flags=re.IGNORECASE)

def generate_individual_exam(room_data):
    num_mcq = min(room_data["active_exam"]["num_mcq"], len(room_data["raw_bank"]["mcq"]))
    num_tf = min(room_data["active_exam"]["num_tf"], len(room_data["raw_bank"]["tf"]))
    num_sa = min(room_data["active_exam"].get("num_sa", 0), len(room_data["raw_bank"].get("sa", [])))

    selected_mcq_raw = random.sample(room_data["raw_bank"]["mcq"], num_mcq) if num_mcq > 0 else []
    selected_tf_raw = random.sample(room_data["raw_bank"]["tf"], num_tf) if num_tf > 0 else []
    selected_sa_raw = random.sample(room_data["raw_bank"]["sa"], num_sa) if num_sa > 0 else []

    random.shuffle(selected_mcq_raw)
    random.shuffle(selected_tf_raw)
    random.shuffle(selected_sa_raw)

    # 1. Trắc nghiệm (G1)
    mcq_send, mcq_server = [], []
    labels = ['A', 'B', 'C', 'D']
    for idx, q in enumerate(selected_mcq_raw, 1):
        shuffled_opts = copy.deepcopy(q['options'])
        random.shuffle(shuffled_opts)
        client_opts, server_opts = [], []
        for o_idx, opt in enumerate(shuffled_opts):
            key = labels[o_idx] if o_idx < len(labels) else opt['key']
            client_opts.append({"key": key, "text": opt["text"]})
            server_opts.append({"key": key, "text": opt["text"], "correct": opt.get("correct", False)})

        clean_q = clean_question_text(q["question"])
        mcq_send.append({"id": idx, "type": "mcq", "question": clean_q, "options": client_opts})
        mcq_server.append({"id": idx, "type": "mcq", "question": clean_q, "options": server_opts})

    # 2. Đúng / Sai (G2)
    tf_send, tf_server = [], []
    for idx, q in enumerate(selected_tf_raw, 1):
        client_subs = [{"key": s["key"], "text": s["text"]} for s in q["sub_items"]]
        server_subs = [{"key": s["key"], "text": s["text"], "correct": s.get("correct", False)} for s in q["sub_items"]]
        clean_q_tf = clean_question_text(q["question"])
        tf_send.append({"id": idx, "type": "tf", "question": clean_q_tf, "sub_items": client_subs})
        tf_server.append({"id": idx, "type": "tf", "question": clean_q_tf, "sub_items": server_subs})

    # 3. Trả lời ngắn (G3)
    sa_send, sa_server = [], []
    for idx, q in enumerate(selected_sa_raw, 1):
        clean_q_sa = clean_question_text(q["question"])
        sa_send.append({"id": idx, "type": "sa", "question": clean_q_sa})
        sa_server.append({"id": idx, "type": "sa", "question": clean_q_sa, "answers": q.get("answers", [])})

    client_payload = {
        "mcq": mcq_send,
        "tf": tf_send,
        "sa": sa_send,
        "duration": room_data["active_exam"]["duration"]
    }
    server_exam = {
        "mcq": mcq_server,
        "tf": tf_server,
        "sa": sa_server
    }
    return client_payload, server_exam

def update_summary_excel(room_id):
    room_data = get_or_create_room(room_id)
    room_dir = get_submission_dir(room_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bang_Diem_Tong_Hop"

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    navy_fill = PatternFill(start_color="0B3C6D", end_color="0B3C6D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)

    headers = ["TT", "Địa chỉ IP", "Họ và tên", "Thời gian còn lại", "Vi phạm", "Trạng thái", "Điểm số"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    scores_list = []
    idx = 1
    current_row = 2

    for sid, st in room_data["students"].items():
        mins = st['time_left'] // 60
        secs = st['time_left'] % 60
        score_val = float(st.get("score", 0.0))
        scores_list.append(score_val)

        ws.append([
            idx,
            st["ip"],
            st["name"],
            f"{mins:02d}:{secs:02d}",
            st["violations"],
            st["submitted"],
            score_val
        ])

        for col_idx in range(1, 8):
            c = ws.cell(row=current_row, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if col_idx == 3:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 7:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = bold_font
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

        idx += 1
        current_row += 1

    total_students = len(scores_list)
    cnt_duoi_5 = sum(1 for s in scores_list if s <= 5.0)
    cnt_duoi_65 = sum(1 for s in scores_list if 5.0 < s <= 6.5)
    cnt_duoi_8 = sum(1 for s in scores_list if 6.5 < s <= 8.0)
    cnt_tren_8 = sum(1 for s in scores_list if s > 8.0)

    def calc_rate(count):
        if total_students == 0:
            return "0.0%"
        return f"{round((count / total_students) * 100, 1)}%"

    stat_start_row = current_row + 2
    stat_headers = ["Mức điểm", "Tỉ lệ", "Số lượng"]
    stat_data = [
        ["<= 5.0", calc_rate(cnt_duoi_5), cnt_duoi_5],
        ["<= 6.5", calc_rate(cnt_duoi_65), cnt_duoi_65],
        ["<= 8.0", calc_rate(cnt_duoi_8), cnt_duoi_8],
        ["> 8", calc_rate(cnt_tren_8), cnt_tren_8]
    ]

    for c_idx, h_text in enumerate(stat_headers, 1):
        cell = ws.cell(row=stat_start_row, column=c_idx, value=h_text)
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for r_offset, row_values in enumerate(stat_data, 1):
        r_idx = stat_start_row + r_offset
        for c_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = {1: 12, 2: 20, 3: 32, 4: 20, 5: 14, 6: 16, 7: 14}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    summary_path = os.path.join(room_dir, "Bang_Diem_Tong_Hop.xlsx")
    wb.save(summary_path)

def save_individual_submission(room_id, st, answers, score):
    room_dir = get_submission_dir(room_id)
    safe_name = re.sub(r'[^\w\s-]', '', st['name']).strip().replace(" ", "_")
    safe_ip = st["ip"].replace(":", "_")
    filename = f"{safe_name}_{safe_ip}.txt"
    filepath = os.path.join(room_dir, filename)

    room_data = get_or_create_room(room_id)
    exam_data = st.get("exam", room_data["raw_bank"])
    NL = "\r\n"

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f" Họ và tên thí sinh : {st['name']}" + NL)
        f.write(f" Địa chỉ IP : {st['ip']}" + NL)
        f.write(f" Trạng thái : {st['submitted']}" + NL)
        f.write(f" Số lần vi phạm : {st['violations']}" + NL)
        f.write(f"Tổng điểm : {score} điểm" + NL)
  
        # PHẦN 1
        f.write("PHẦN 1: TRẮC NGHIỆM NHIỀU LỰA CHỌN" + NL)
        f.write("--------------------------------------------------------" + NL + NL)
        mcq_list = exam_data.get("mcq", [])
        for idx, q in enumerate(mcq_list, 1):
            ans = answers.get(f"mcq_{q['id']}", "Chưa làm")
            correct_opt = next((opt['key'] for opt in q['options'] if opt.get('correct')), "Chưa rõ")
            f.write(f"Câu {idx}: {q['question']}" + NL)
            for opt in q['options']:
                mark = "[ĐÚNG]" if opt.get('correct') else "      "
                f.write(f"   {mark} {opt['key']}. {opt['text']}" + NL)
            f.write(f"   Thí sinh : {ans}" + NL)
            f.write(f"   Đáp án : {correct_opt}" + NL)
            f.write(NL)

        # PHẦN 2
        f.write(" PHẦN 2: TRẮC NGHIỆM ĐÚNG / SAI" + NL)
        f.write("---------------------------------------------" + NL + NL)
        tf_list = exam_data.get("tf", [])
        for idx, q in enumerate(tf_list, 1):
            f.write(f"Câu {idx}: {q['question']}" + NL)
            for sub in q.get('sub_items', []):
                user_val = answers.get(f"tf_{q['id']}_{sub['key']}", "Chưa làm")
                expected = "Đ" if sub.get('correct') else "S"
                f.write(f"   {sub['key']}) {sub['text']}" + NL)
                f.write(f"   Thí sinh : {user_val}" + NL)
                f.write(f"   Đáp án : {expected}" + NL)
            f.write(NL)

        # PHẦN 3
        f.write(" PHẦN 3: CÂU HỎI TRẢ LỜI NGẮN" + NL)
        f.write("---------------------------------------------" + NL + NL)
        sa_list = exam_data.get("sa", [])
        for idx, q in enumerate(sa_list, 1):
            user_val = answers.get(f"sa_{q['id']}", "Chưa làm")
            expected_list = q.get("answers", [])
            expected_str = " / ".join(expected_list) if expected_list else "Chưa rõ"
            f.write(f"Câu {idx}: {q['question']}" + NL)
            f.write(f"   Thí sinh : {user_val}" + NL)
            f.write(f"   Đáp án chuẩn : {expected_str}" + NL)
            f.write(NL)

def process_grading(room_id, sid, answers=None, student_name=None):
    room_data = get_or_create_room(room_id)
    st = None
    target_sid = sid

    if sid in room_data["students"]:
        st = room_data["students"][sid]
    elif student_name:
        for s_id, s_data in room_data["students"].items():
            if s_data.get("name") == student_name and s_data.get("submitted") != "Đã nộp":
                st = s_data
                target_sid = s_id
                break

    if not st:
        return

    if st.get("submitted") == "Đã nộp":
        socketio.emit('exam_finished_ack', {"score": st.get("score", 0.0)}, room=sid)
        if target_sid != sid:
            socketio.emit('exam_finished_ack', {"score": st.get("score", 0.0)}, room=target_sid)
        return

    if not answers and st.get("answers"):
        answers = st["answers"]
    elif not answers:
        answers = {}

    st["submitted"] = "Đã nộp"
    st["answers"] = answers
    st["time_left"] = 0
    room_data["submitted_ips"].add(st["ip"])

    student_exam = st.get("exam")
    if not student_exam:
        student_exam = {
            "mcq": room_data["raw_bank"]["mcq"][:room_data["active_exam"]["num_mcq"]],
            "tf": room_data["raw_bank"]["tf"][:room_data["active_exam"]["num_tf"]],
            "sa": room_data["raw_bank"].get("sa", [])[:room_data["active_exam"].get("num_sa", 0)]
        }

    score = 0.0

    # Chấm MCQ
    mcq_pts = room_data["active_exam"].get("score_per_mcq", 0.25)
    for q in student_exam.get("mcq", []):
        ans = answers.get(f"mcq_{q['id']}")
        correct_opt = next((opt['key'] for opt in q['options'] if opt.get('correct')), None)
        if ans and str(ans).strip().upper() == str(correct_opt).strip().upper():
            score += mcq_pts

    # Chấm TF
    tf_pts = room_data["active_exam"].get("score_per_tf", 1.0)
    scale = room_data["active_exam"].get("tf_scale", {"1": 0.1, "2": 0.25, "3": 0.5, "4": 1.0})
    for q in student_exam.get("tf", []):
        correct_cnt = 0
        for sub in q.get('sub_items', []):
            user_val = answers.get(f"tf_{q['id']}_{sub['key']}")
            expected = "Đ" if sub.get('correct') else "S"
            if user_val and str(user_val).strip() == expected:
                correct_cnt += 1
        ratio = float(scale.get(str(correct_cnt), 0.0))
        score += tf_pts * ratio

    # Chấm SA
    sa_pts = room_data["active_exam"].get("score_per_sa", 0.0)
    for q in student_exam.get("sa", []):
        user_val = answers.get(f"sa_{q['id']}", "")
        if user_val:
            cleaned_user = str(user_val).strip().lower()
            valid_answers = [str(a).strip().lower() for a in q.get("answers", [])]
            if cleaned_user in valid_answers:
                score += sa_pts

    st["score"] = round(score, 2)

    save_individual_submission(room_id, st, answers, st["score"])
    update_summary_excel(room_id)

    socketio.emit('update_teacher_list', list(room_data["students"].values()), to=room_id)
    socketio.emit('exam_finished_ack', {"score": st["score"]}, room=sid)
    if target_sid != sid:
        socketio.emit('exam_finished_ack', {"score": st["score"]}, room=target_sid)

@app.route('/')
def teacher_dashboard():
    return render_template('index.html')

@app.route('/student')
def student_view():
    return render_template('student.html')

# ROUTE RÚT GỌN LINK CHO HỌC SINH (/r/mã_phòng)
@app.route('/r/<room_id>')
def short_room(room_id):
    return redirect(f"/student?room={room_id}")

# ROUTE SIÊU NHẸ ĐỂ PING CHỐNG NGỦ ĐÔNG
@app.route('/ping')
def ping():
    return "pong", 200

@app.route('/api/get_ip', methods=['GET'])
def api_get_ip():
    room_id = request.args.get('room', '')
    base_url = get_base_url()
    # Mặc định trả về link ngắn /r/room_id
    if room_id:
        url = f"{base_url}/r/{room_id}"
    else:
        url = f"{base_url}/student"
    return jsonify({
        "ip": base_url.replace("http://", "").replace("https://", "").split(":")[0],
        "port": 5000,
        "url": url
    })

@app.route('/api/upload_bank', methods=['POST'])
def upload_bank():
    room_id = request.form.get('room', 'default')
    room_data = get_or_create_room(room_id)

    if 'file' not in request.files:
        return jsonify({"success": False, "error": "Không tìm thấy file tải lên!"})
    
    file = request.files['file']
    filename = file.filename
    save_path = os.path.join(UPLOAD_DIR, f"{room_id}_{filename}")
    file.save(save_path)

    paragraphs = extract_docx_with_cloudinary(save_path)

    mcq_questions = []
    tf_questions = []
    sa_questions = []
    current_mode = None
    curr_q = None
    q_counter_mcq = 0
    q_counter_tf = 0
    q_counter_sa = 0

    for line in paragraphs:
        lower_line = line.lower()
        if lower_line.startswith('g1') or 'phần 1' in lower_line or 'phan 1' in lower_line or 'part 1' in lower_line:
            if curr_q:
                if current_mode == 'g1': mcq_questions.append(curr_q)
                elif current_mode == 'g2': tf_questions.append(curr_q)
                elif current_mode == 'g3': sa_questions.append(curr_q)
                curr_q = None
            current_mode = 'g1'
            continue
        elif lower_line.startswith('g2') or 'phần 2' in lower_line or 'phan 2' in lower_line or 'part 2' in lower_line:
            if curr_q:
                if current_mode == 'g1': mcq_questions.append(curr_q)
                elif current_mode == 'g2': tf_questions.append(curr_q)
                elif current_mode == 'g3': sa_questions.append(curr_q)
                curr_q = None
            current_mode = 'g2'
            continue
        elif lower_line.startswith('g3') or 'phần 3' in lower_line or 'phan 3' in lower_line or 'part 3' in lower_line:
            if curr_q:
                if current_mode == 'g1': mcq_questions.append(curr_q)
                elif current_mode == 'g2': tf_questions.append(curr_q)
                elif current_mode == 'g3': sa_questions.append(curr_q)
                curr_q = None
            current_mode = 'g3'
            continue

        if current_mode == 'g1':
            match_opt = re.match(r'^(#?)\s*([A-Da-d])[\.\)]\s*(.*)', line, re.DOTALL)
            if match_opt:
                if not curr_q:
                    return jsonify({"success": False, "error": f"Lỗi tại phương án: '{line}'. Chưa có câu hỏi!"})
                is_correct = bool(match_opt.group(1) == '#')
                curr_q['options'].append({"key": match_opt.group(2).upper(), "text": match_opt.group(3), "correct": is_correct})
            else:
                is_new_q = re.match(r'^câu\s*\d+', lower_line) or (curr_q is None)
                if is_new_q:
                    if curr_q: mcq_questions.append(curr_q)
                    q_counter_mcq += 1
                    curr_q = {"id": q_counter_mcq, "type": "mcq", "question": line, "options": []}
                else:
                    curr_q["question"] += "<br>" + line

        elif current_mode == 'g2':
            match_tf_opt = re.match(r'^(\*?)\s*([a-dA-D])[\.\)]\s*(.*)', line, re.DOTALL)
            if match_tf_opt:
                if not curr_q:
                    return jsonify({"success": False, "error": f"Lỗi tại ý Đ/S: '{line}'. Chưa có câu hỏi!"})
                is_correct = bool(match_tf_opt.group(1) == '*')
                curr_q['sub_items'].append({"key": match_tf_opt.group(2).lower(), "text": match_tf_opt.group(3), "correct": is_correct})
            else:
                is_new_q = re.match(r'^câu\s*\d+', lower_line) or (curr_q is None)
                if is_new_q:
                    if curr_q: tf_questions.append(curr_q)
                    q_counter_tf += 1
                    curr_q = {"id": q_counter_tf, "type": "tf", "question": line, "sub_items": []}
                else:
                    curr_q["question"] += "<br>" + line

        elif current_mode == 'g3':
            match_sa_ans = re.match(r'^(?:#|đáp\s*án\s*[\:\.]?\s*)(.+)', line, re.IGNORECASE)
            if match_sa_ans:
                if not curr_q:
                    return jsonify({"success": False, "error": f"Lỗi tại đáp án ngắn: '{line}'. Chưa có câu hỏi!"})
                raw_ans_list = match_sa_ans.group(1).split('|')
                curr_q['answers'] = [a.strip() for a in raw_ans_list if a.strip()]
            else:
                is_new_q = re.match(r'^câu\s*\d+', lower_line) or (curr_q is None)
                if is_new_q:
                    if curr_q: sa_questions.append(curr_q)
                    q_counter_sa += 1
                    curr_q = {"id": q_counter_sa, "type": "sa", "question": line, "answers": []}
                else:
                    curr_q["question"] += "<br>" + line

    if curr_q:
        if current_mode == 'g1': mcq_questions.append(curr_q)
        elif current_mode == 'g2': tf_questions.append(curr_q)
        elif current_mode == 'g3': sa_questions.append(curr_q)

    for q in mcq_questions:
        num_c = sum(1 for opt in q['options'] if opt['correct'])
        if len(q['options']) < 2 or num_c != 1:
            return jsonify({"success": False, "error": f"Câu {q['id']} (Phần 1) phải có các phương án và đúng 1 dấu (#)!"})

    for q in tf_questions:
        if len(q['sub_items']) == 0:
            return jsonify({"success": False, "error": f"Câu {q['id']} (Phần 2) chưa có các ý hỏi con!"})

    for q in sa_questions:
        if len(q.get('answers', [])) == 0:
            return jsonify({"success": False, "error": f"Câu {q['id']} (Phần 3) chưa có đáp án đúng (bắt đầu bằng dấu #)!"})

    room_data["raw_bank"] = {
        "filename": filename,
        "mcq": mcq_questions,
        "tf": tf_questions,
        "sa": sa_questions
    }
    return jsonify({
        "success": True,
        "filename": filename,
        "total_mcq": len(mcq_questions),
        "total_tf": len(tf_questions),
        "total_sa": len(sa_questions)
    })

@app.route('/api/configure_exam', methods=['POST'])
def configure_exam():
    data = request.json or {}
    room_id = data.get("room", "default")
    room_data = get_or_create_room(room_id)

    room_data["submitted_ips"].clear()
    room_data["students"].clear()
    clear_submission_folder(room_id)

    n_mcq = data.get("num_mcq", 0)
    s_mcq = data.get("score_mcq", 0.0)
    n_tf = data.get("num_tf", 0)
    s_tf = data.get("score_tf", 0.0)
    n_sa = data.get("num_sa", 0)
    s_sa = data.get("score_sa", 0.0)
    dur = data.get("duration", 15)

    room_data["class_name"] = data.get("class_name", "")
    room_data["active_exam"].update({
        "num_mcq": n_mcq,
        "score_mcq": s_mcq,
        "score_per_mcq": round(s_mcq / max(1, n_mcq), 3),
        "num_tf": n_tf,
        "score_tf": s_tf,
        "score_per_tf": round(s_tf / max(1, n_tf), 3),
        "num_sa": n_sa,
        "score_sa": s_sa,
        "score_per_sa": round(s_sa / max(1, n_sa), 3),
        "duration": dur,
        "tf_scale": data.get("tf_scale", {"1": 0.1, "2": 0.25, "3": 0.5, "4": 1.0}),
        "status": "ready"
    })
    socketio.emit('update_teacher_list', [], to=room_id)
    return jsonify({"success": True})

@app.route('/api/reset_exam', methods=['POST'])
def api_reset_exam():
    data = request.json or {}
    room_id = data.get("room", "default")
    room_data = get_or_create_room(room_id)

    room_data["submitted_ips"].clear()
    room_data["students"].clear()
    clear_submission_folder(room_id)
    room_data["class_name"] = ""
    room_data["active_exam"] = {
        "num_mcq": 0, "score_mcq": 0.0, "score_per_mcq": 0.0,
        "num_tf": 0, "score_tf": 0.0, "score_per_tf": 0.0,
        "num_sa": 0, "score_sa": 0.0, "score_per_sa": 0.0,
        "duration": 15, "tf_scale": {"1": 0.1, "2": 0.25, "3": 0.5, "4": 1.0},
        "status": "waiting", "start_time": None
    }
    socketio.emit('update_teacher_list', [], to=room_id)
    return jsonify({"success": True})

@app.route('/api/export_excel', methods=['GET'])
def export_excel():
    room_id = request.args.get('room', 'default')
    room_dir = get_submission_dir(room_id)
    summary_path = os.path.join(room_dir, "Bang_Diem_Tong_Hop.xlsx")
    if os.path.exists(summary_path):
        return send_file(summary_path, as_attachment=True, download_name="Bang_Diem_Tong_Hop.xlsx")
    update_summary_excel(room_id)
    return send_file(summary_path, as_attachment=True, download_name="Bang_Diem_Tong_Hop.xlsx")

@app.route('/api/download_all_results', methods=['GET'])
def download_all_results():
    room_id = request.args.get('room', 'default')
    room_data = get_or_create_room(room_id)
    room_dir = get_submission_dir(room_id)

    update_summary_excel(room_id)
    memory_file = BytesIO()
    with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(room_dir):
            for file in files:
                file_path = os.path.join(root, file)
                zf.write(file_path, arcname=file)

    memory_file.seek(0)
    class_name = room_data.get("class_name", f"Phong_{room_id}").replace(" ", "_")
    zip_filename = f"KetQua_BaiThi_{class_name}.zip"

    return send_file(
        memory_file,
        mimetype='application/zip',
        as_attachment=True,
        download_name=zip_filename
    )

@socketio.on('join_teacher')
def handle_teacher_join(data):
    room_id = data.get('room', 'default')
    join_room(room_id)
    room_data = get_or_create_room(room_id)
    emit('update_teacher_list', list(room_data["students"].values()), room=request.sid)

@socketio.on('join_student')
def handle_student_join(data):
    sid = request.sid
    room_id = data.get('room', 'default')
    join_room(room_id)
    room_data = get_or_create_room(room_id)
    client_ip = get_client_ip(request)

    if client_ip in room_data["submitted_ips"]:
        emit('exam_blocked', {'message': 'Bạn đã thi rồi trong phòng này!'}, room=sid)
        return

    name = data.get('name', 'Thí sinh').strip() or 'Thí sinh'

    existing_sid = None
    for s_id, s_data in room_data["students"].items():
        if s_data["name"] == name and s_data["ip"] == client_ip:
            existing_sid = s_id
            break

    time_left = room_data["active_exam"]["duration"] * 60
    if room_data["active_exam"]["status"] == "running" and room_data["active_exam"]["start_time"]:
        elapsed = int(time.time() - room_data["active_exam"]["start_time"])
        time_left = max(0, time_left - elapsed)

    if existing_sid:
        student_record = room_data["students"].pop(existing_sid)
        student_record["time_left"] = time_left
        room_data["students"][sid] = student_record
    else:
        room_data["students"][sid] = {
            "ip": client_ip,
            "name": name,
            "room": room_id,
            "time_left": time_left,
            "violations": 0,
            "submitted": "Đang thi",
            "score": 0.0,
            "answers": {},
            "exam": None
        }

    emit('update_teacher_list', list(room_data["students"].values()), to=room_id)

    if room_data["active_exam"]["status"] == "running":
        if not room_data["students"][sid].get("exam"):
            client_payload, server_exam = generate_individual_exam(room_data)
            room_data["students"][sid]["exam"] = server_exam
            client_payload["time_left"] = time_left
            emit('start_exam_now', client_payload, room=sid)
        else:
            client_payload = {
                "mcq": [{"id": q["id"], "type": "mcq", "question": q["question"], "options": [{"key": o["key"], "text": o["text"]} for o in q["options"]]} for q in room_data["students"][sid]["exam"]["mcq"]],
                "tf": [{"id": q["id"], "type": "tf", "question": q["question"], "sub_items": [{"key": s["key"], "text": s["text"]} for s in q["sub_items"]]} for q in room_data["students"][sid]["exam"]["tf"]],
                "sa": [{"id": q["id"], "type": "sa", "question": q["question"]} for q in room_data["students"][sid]["exam"].get("sa", [])],
                "duration": room_data["active_exam"]["duration"],
                "time_left": time_left
            }
            emit('start_exam_now', client_payload, room=sid)

@socketio.on('sync_answer')
def handle_sync_answer(data):
    sid = request.sid
    room_id = data.get('room', 'default')
    room_data = get_or_create_room(room_id)
    if sid in room_data["students"]:
        key = data.get("key")
        val = data.get("val")
        if key:
            room_data["students"][sid]["answers"][key] = val

@socketio.on('sync_timer')
def handle_timer_sync(data):
    sid = request.sid
    room_id = data.get('room', 'default')
    room_data = get_or_create_room(room_id)
    if sid in room_data["students"]:
        st = room_data["students"][sid]
        t_left = data.get("time_left", 0)
        st["time_left"] = t_left

        if t_left <= 0 and st["submitted"] != "Đã nộp":
            process_grading(room_id, sid, st.get("answers", {}), st.get("name"))

@socketio.on('report_violation')
def handle_violation(data):
    sid = request.sid
    room_id = data.get('room', 'default') if isinstance(data, dict) else 'default'
    room_data = get_or_create_room(room_id)
    if sid in room_data["students"]:
        room_data["students"][sid]["violations"] += 1
        emit('update_teacher_list', list(room_data["students"].values()), to=room_id)

@socketio.on('submit_exam')
def handle_submit(data):
    sid = request.sid
    room_id = data.get('room', 'default') if isinstance(data, dict) else 'default'
    if isinstance(data, dict) and 'answers' in data:
        process_grading(room_id, sid, data.get('answers', {}), data.get('name'))
    else:
        process_grading(room_id, sid, data if isinstance(data, dict) else {})

@socketio.on('teacher_start_exam')
def teacher_start(data):
    room_id = data.get('room', 'default') if isinstance(data, dict) else 'default'
    room_data = get_or_create_room(room_id)

    room_data["submitted_ips"].clear()
    clear_submission_folder(room_id)

    room_data["active_exam"]["status"] = "running"
    room_data["active_exam"]["start_time"] = time.time()
    
    duration_secs = room_data["active_exam"]["duration"] * 60
    for sid, st in room_data["students"].items():
        client_payload, server_exam = generate_individual_exam(room_data)
        st["exam"] = server_exam
        client_payload["time_left"] = duration_secs
        emit('start_exam_now', client_payload, room=sid)

@socketio.on('reset_exam_session')
def handle_reset_session(data):
    room_id = data.get('room', 'default') if isinstance(data, dict) else 'default'
    room_data = get_or_create_room(room_id)

    room_data["submitted_ips"].clear()
    room_data["students"].clear()
    clear_submission_folder(room_id)
    room_data["active_exam"]["status"] = "waiting"
    room_data["active_exam"]["start_time"] = None
    emit('update_teacher_list', [], to=room_id)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    socketio.run(app, host='0.0.0.0', port=port, debug=False)
